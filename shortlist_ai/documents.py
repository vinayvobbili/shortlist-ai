"""Turn resume / job-description files into text (and, for PDFs, a document block
that vision-capable backends can read directly)."""

import base64
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PageObject, PdfReader
from pypdf.generic import ContentStream, DictionaryObject, NameObject

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class LoadedDocument:
    path: Path
    text: str                 # visible text; empty for scanned PDFs
    pdf_base64: str | None    # set for PDFs, so backends that read pages can use them
    hidden_text: str = ""     # PDF text a reader can't see (white or tiny): a prompt-injection vector

    def content_blocks(self) -> list[dict]:
        """Content for a backend that accepts documents: the PDF itself when available."""
        if self.pdf_base64:
            return [{"type": "document",
                     "source": {"type": "base64", "media_type": "application/pdf", "data": self.pdf_base64}}]
        return [{"type": "text", "text": self.text}]


# Text below this effective size (points) is unreadable when printed; resumes don't use it legitimately.
MIN_VISIBLE_PT = 4.0
# Text whose colour is this close to what's behind it (max channel difference, 0-1) can't be read.
MIN_CONTRAST = 0.12
# Fill opacity below which text can't be read, and below which a filled shape doesn't hide the page.
MIN_TEXT_OPACITY, MIN_BACKGROUND_OPACITY = 0.15, 0.5

_SHOW_TEXT = {b"Tj", b"TJ", b"'", b'"'}
_FILL = {b"f", b"F", b"f*", b"B", b"B*", b"b", b"b*"}
_END_PATH = {b"n", b"S", b"s"}
Matrix = tuple[float, float, float, float, float, float]
_IDENTITY: Matrix = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


def _mul(m: Matrix, n: Matrix) -> Matrix:
    a, b, c, d, e, f = m
    a2, b2, c2, d2, e2, f2 = n
    return (a * a2 + b * c2, a * b2 + b * d2, c * a2 + d * c2, c * b2 + d * d2,
            e * a2 + f * c2 + e2, e * b2 + f * d2 + f2)


def _apply(m: Matrix, x: float, y: float) -> tuple[float, float]:
    return x * m[0] + y * m[2] + m[4], x * m[1] + y * m[3] + m[5]


def _color(args) -> tuple[float, float, float] | None:
    """Gray, RGB or CMYK operands as RGB; None when it can't be known (patterns, named colours)."""
    try:
        v = [float(a) for a in args]
    except (TypeError, ValueError):
        return None
    if len(v) == 1:
        return (v[0],) * 3
    if len(v) == 3:
        return tuple(v)
    if len(v) == 4:
        c, m, y, k = v
        return tuple((1 - x) * (1 - k) for x in (c, m, y))
    return None


def _bbox(points) -> tuple[float, float, float, float]:
    return min(p[0] for p in points), min(p[1] for p in points), max(p[0] for p in points), max(p[1] for p in points)


def _intersect(a, b):
    """Intersection of two boxes, or None when they don't overlap."""
    if a is None or b is None:
        return None
    box = (max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3]))
    return box if box[0] <= box[2] and box[1] <= box[3] else None


def _solid_boxes(subpaths) -> list[tuple[float, float, float, float]]:
    """Bounding boxes of a filled path's subpaths. Subpaths that nest are likely a frame and its
    hole (even-odd or opposite winding), so neither is treated as a solid background."""
    boxes = [_bbox(sp) for sp in subpaths if sp]

    def encloses(a, b):
        return a is not b and a[0] <= b[0] and a[1] <= b[1] and a[2] >= b[2] and a[3] >= b[3]

    return [a for a in boxes if not any(encloses(a, b) or encloses(b, a) for b in boxes)]


class _PageScan:
    """Replays a page's content stream and classifies each text-drawing operation as visible
    or hidden, using the state at the moment it is drawn."""

    def __init__(self, page, operations):
        self.page = page
        box = page.mediabox
        self.bounds = (float(box.left) - 1, float(box.bottom) - 1, float(box.right) + 1, float(box.top) + 1)
        self.ops = operations
        self.hidden_ops: set[int] = set()
        # Painted regions in paint order: (x0, y0, x1, y1, colour); colour None = image or unknown.
        self.regions: list[tuple[float, float, float, float, tuple | None]] = []
        self.has_image = any(op == b"INLINE IMAGE" or (op == b"Do" and self._is_image(args))
                             for args, op in operations)

    def _opacity(self, args, current: float) -> float:
        try:
            return float(self.page["/Resources"]["/ExtGState"][args[0]].get_object().get("/ca", current))
        except (KeyError, TypeError, IndexError, ValueError):
            return current

    def _xobject(self, args):
        try:
            return self.page["/Resources"]["/XObject"][args[0]].get_object()
        except (KeyError, TypeError, IndexError):
            return None

    def _is_image(self, args) -> bool:
        xobj = self._xobject(args)
        return xobj is not None and xobj.get("/Subtype") == "/Image"

    def _paint_unknown(self, box):
        """Something whose colour we don't track (image, gradient, form) was painted over box."""
        if box:
            self.regions.append((*box, None))

    def background(self, x: float, y: float):
        for x0, y0, x1, y1, color in reversed(self.regions):
            if x0 <= x <= x1 and y0 <= y <= y1:
                return color
        return (1.0, 1.0, 1.0)  # the page itself

    def run(self) -> "_PageScan":
        ctm, fill, render, size, leading, alpha = _IDENTITY, (0.0, 0.0, 0.0), 0, 0.0, 0.0, 1.0
        clip = self.bounds          # bounding box of the clipping path, device space
        clipping = False            # a W/W* is waiting for the path to end
        stack = []
        tlm = tm = _IDENTITY
        path: list[list[tuple[float, float]]] = []  # subpaths, each a list of device-space points

        def next_line(tx, ty):
            nonlocal tlm, tm
            tlm = _mul((1.0, 0.0, 0.0, 1.0, tx, ty), tlm)
            tm = tlm

        for i, (args, op) in enumerate(self.ops):
            try:
                if op == b"q":
                    stack.append((ctm, fill, render, size, leading, alpha, clip))
                elif op == b"Q" and stack:
                    ctm, fill, render, size, leading, alpha, clip = stack.pop()
                elif op in (b"W", b"W*"):
                    clipping = True
                elif op == b"gs":
                    alpha = self._opacity(args, alpha)
                elif op == b"cm":
                    ctm = _mul(tuple(float(a) for a in args), ctm)
                elif op in (b"rg", b"g", b"k", b"sc", b"scn"):
                    fill = _color(args)
                elif op == b"cs":
                    fill = (0.0, 0.0, 0.0)  # initial colour of the new space
                elif op == b"re":
                    x, y, w, h = (float(a) for a in args)
                    path.append([_apply(ctm, x, y), _apply(ctm, x + w, y + h)])
                elif op in (b"m", b"l", b"c", b"v", b"y"):
                    nums = [float(a) for a in args]
                    points = [_apply(ctm, nums[j], nums[j + 1]) for j in range(0, len(nums), 2)]
                    if op == b"m" or not path:
                        path.append(points)
                    else:
                        path[-1] += points
                elif op in _FILL or op in _END_PATH:
                    if op in _FILL and alpha >= MIN_BACKGROUND_OPACITY:
                        self.regions += [(*box, fill) for box in _solid_boxes(path)]
                    if clipping and path:
                        clip = _intersect(clip, _bbox([p for sp in path for p in sp]))
                    path, clipping = [], False
                elif op == b"sh":
                    self._paint_unknown(clip)
                elif op == b"INLINE IMAGE" or (op == b"Do" and self._is_image(args)):
                    self._paint_unknown(_intersect(clip, _bbox([_apply(ctm, x, y) for x in (0, 1) for y in (0, 1)])))
                elif op == b"Do" and (form := self._xobject(args)) is not None and "/BBox" in form:
                    x0, y0, x1, y1 = (float(v) for v in form["/BBox"])
                    corners = [_apply(ctm, x, y) for x in (x0, x1) for y in (y0, y1)]
                    self._paint_unknown(_intersect(clip, _bbox(corners)))
                elif op == b"BT":
                    tlm = tm = _IDENTITY
                elif op == b"Tm":
                    tlm = tm = tuple(float(a) for a in args)
                elif op == b"Td":
                    next_line(float(args[0]), float(args[1]))
                elif op == b"TD":
                    leading = -float(args[1])
                    next_line(float(args[0]), float(args[1]))
                elif op == b"TL":
                    leading = float(args[0])
                elif op == b"T*":
                    next_line(0.0, -leading)
                elif op == b"Tf":
                    size = float(args[1])
                elif op == b"Tr":
                    render = int(args[0])
                if op in (b"'", b'"'):
                    next_line(0.0, -leading)
                if op in _SHOW_TEXT and self._hidden(_mul(tm, ctm), fill, render, size, alpha):
                    self.hidden_ops.add(i)
            except (TypeError, ValueError, IndexError, ZeroDivisionError):
                continue  # malformed operator: skip it rather than fail the whole resume
        return self

    def _hidden(self, m: Matrix, fill, render: int, size: float, alpha: float) -> bool:
        if render in (3, 7) and not self.has_image:
            return True  # invisible text is legitimate only as an OCR layer over a scan
        if alpha < MIN_TEXT_OPACITY and render not in (3, 7):
            return True
        if size and size * (m[2] ** 2 + m[3] ** 2) ** 0.5 < MIN_VISIBLE_PT:
            return True
        x, y = m[4], m[5]
        x0, y0, x1, y1 = self.bounds
        if not (x0 <= x <= x1 and y0 <= y <= y1):
            return True  # drawn off the page
        behind = self.background(x, y)
        if fill is None or behind is None:
            return False  # can't tell (pattern colour, or text over an image): assume visible
        return max(abs(a - b) for a, b in zip(fill, behind, strict=True)) < MIN_CONTRAST


def _pdf_text(path: Path) -> tuple[str, str]:
    """Split a PDF's text layer into what a reader sees and what they don't.

    Hidden = text that can't be seen where it is drawn: the same colour as what's behind it
    (white on white, but not white on a dark sidebar), transparent, below MIN_VISIBLE_PT, off the
    page, or in an invisible render mode on a page without a scanned image. These are the usual
    ways to plant instructions or keywords that only a parser reads. Text over something whose colour isn't
    tracked (images, gradients, form XObjects) is assumed visible, and text inside form XObjects
    isn't classified; skills grounding still applies to all of these.
    """
    reader = PdfReader(path)
    visible, hidden = [], []
    for page in reader.pages:
        contents = page.get_contents()
        if contents is None:
            continue
        operations = ContentStream(contents, reader).operations
        scan = _PageScan(page, operations).run()
        if not scan.hidden_ops:
            visible.append(page.extract_text() or "")
            continue
        shown = [op for i, op in enumerate(operations) if i not in scan.hidden_ops]
        # The hidden copy drops visible text and form XObjects (their text counts as visible).
        only_hidden = [op for i, op in enumerate(operations)
                       if i in scan.hidden_ops or (op[1] not in _SHOW_TEXT and op[1] != b"Do")]
        visible.append(_extract(reader, page, shown))
        hidden.append(_extract(reader, page, only_hidden))
    return "\n".join(visible).strip(), " ".join(" ".join(h.split()) for h in hidden if h.strip())


def _extract(reader, page, operations) -> str:
    stream = ContentStream(None, reader)
    stream.operations = operations
    copy = PageObject.create_blank_page(reader, float(page.mediabox.width), float(page.mediabox.height))
    copy[NameObject("/Resources")] = page["/Resources"] if "/Resources" in page else DictionaryObject()
    copy[NameObject("/MediaBox")] = page.mediabox
    copy.replace_contents(stream)
    return copy.extract_text() or ""


def load_document(path: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix} (supported: {', '.join(sorted(SUPPORTED))})")

    if suffix == ".pdf":
        visible, hidden = _pdf_text(path)
        return LoadedDocument(path, visible, base64.standard_b64encode(path.read_bytes()).decode(), hidden)

    if suffix == ".docx":
        doc = DocxDocument(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many resume templates put content in tables
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(p for p in parts if p.strip())
    else:
        text = path.read_text(encoding="utf-8", errors="replace")

    if not text.strip():
        raise ValueError(f"No text found in {path}")
    return LoadedDocument(path, text, None)
